"""Extract the readable text from a directory tree of DOCX references.

This is a mechanical, read-only corpus conversion.  It preserves the relative
folder structure, emits one UTF-8 text file per DOCX, and builds a combined
corpus with explicit source boundaries for downstream literary analysis.
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document


def paragraph_text(paragraph) -> str:
    return "".join(run.text for run in paragraph.runs).strip()


def extract_docx(path: Path) -> str:
    document = Document(path)
    blocks: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph_text(paragraph)
        if text:
            blocks.append(text)

    for table_index, table in enumerate(document.tables, 1):
        rows: list[str] = []
        for row in table.rows:
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            blocks.append(f"[TABLEAU {table_index}]")
            blocks.extend(rows)

    return "\n\n".join(blocks).strip() + ("\n" if blocks else "")


def safe_output_path(source_root: Path, output_root: Path, source: Path) -> Path:
    relative = source.relative_to(source_root)
    return output_root / relative.with_suffix(".txt")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source_root", type=Path)
    parser.add_argument("output_root", type=Path)
    args = parser.parse_args()

    source_root = args.source_root.resolve()
    output_root = args.output_root.resolve()
    output_root.mkdir(parents=True, exist_ok=True)

    manifest: list[dict[str, object]] = []
    combined: list[str] = []

    for source in sorted(source_root.rglob("*.docx"), key=lambda p: str(p).casefold()):
        output = safe_output_path(source_root, output_root, source)
        output.parent.mkdir(parents=True, exist_ok=True)
        try:
            text = extract_docx(source)
            error = None
        except Exception as exc:  # Keep the batch useful if one legacy file is malformed.
            text = ""
            error = f"{type(exc).__name__}: {exc}"

        output.write_text(text, encoding="utf-8")
        words = len(re.findall(r"\b[\wÀ-ÖØ-öø-ÿŒœ’'-]+\b", text, flags=re.UNICODE))
        record = {
            "source": str(source),
            "relative": str(source.relative_to(source_root)),
            "output": str(output),
            "characters": len(text),
            "words": words,
            "error": error,
        }
        manifest.append(record)

        combined.extend(
            [
                "=" * 90,
                f"SOURCE DOCX : {record['relative']}",
                f"MOTS : {words}",
                "=" * 90,
                text.rstrip(),
                "",
            ]
        )

    (output_root / "_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    (output_root / "_corpus_complet.txt").write_text(
        "\n".join(combined).rstrip() + "\n",
        encoding="utf-8",
    )

    summary = {
        "documents": len(manifest),
        "documents_with_errors": sum(1 for item in manifest if item["error"]),
        "characters": sum(int(item["characters"]) for item in manifest),
        "words": sum(int(item["words"]) for item in manifest),
        "output_root": str(output_root),
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
